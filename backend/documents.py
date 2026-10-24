"""Document attachments for contacts and pitching companies.

Uploaded files are never written to disk. Text is extracted in memory and
stored in Neon (or the local JSON fallback) so the LLM can ground its writing
in a CV, bio, company brief, or similar reference material.
"""

from __future__ import annotations

import io
import json
import re
import threading
import uuid
from datetime import datetime, timezone
from pathlib import Path

from . import config, database

MAX_FILE_BYTES = 5 * 1024 * 1024  # 5 MB per file
MAX_TEXT_CHARS = 50_000  # per document
MAX_PROMPT_CHARS = 12_000  # total reference text sent to the LLM
MAX_DOCUMENTS_PER_OWNER = 10
MAX_PDF_PAGES = 100

OWNER_TYPES = {"contact", "company"}

ALLOWED_EXTENSIONS = {
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".txt": "text/plain",
    ".md": "text/markdown",
}

_FILE_LOCK = threading.Lock()


class DocumentError(ValueError):
    """Raised for user-facing document validation or extraction failures."""


def _get_connection():
    return database._connect()


def normalize_owner(owner_type: str, owner_key: str) -> tuple[str, str]:
    owner_type = (owner_type or "").strip().lower()
    if owner_type not in OWNER_TYPES:
        raise DocumentError("owner_type must be either 'contact' or 'company'.")
    key = (owner_key or "").strip()
    if not key:
        raise DocumentError("owner_key is required.")
    if owner_type == "contact":
        key = key.lower()
    if len(key) > 320:
        raise DocumentError("owner_key is too long.")
    return owner_type, key


def _clean_text(text: str) -> str:
    text = (text or "").replace("\x00", "")
    text = "".join(ch for ch in text if ch in "\n\r\t" or ord(ch) >= 32)
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"\n{4,}", "\n\n\n", text)
    return text.strip()[:MAX_TEXT_CHARS]


def sanitize_filename(filename: str) -> str:
    name = Path(filename or "document").name
    name = re.sub(r"[^\w\s.\-()\[\]]", "_", name, flags=re.UNICODE).strip(" .")
    return (name or "document")[:200]


def _extract_pdf(data: bytes) -> str:
    import pymupdf

    try:
        doc = pymupdf.open(stream=data, filetype="pdf")
    except Exception as error:
        raise DocumentError("This PDF could not be opened.") from error

    try:
        if doc.is_encrypted or doc.needs_pass:
            raise DocumentError("Password-protected PDFs are not supported. Upload an unlocked copy.")
        if doc.page_count > MAX_PDF_PAGES:
            raise DocumentError(f"PDFs are limited to {MAX_PDF_PAGES} pages.")

        parts: list[str] = []
        total = 0
        for page in doc:
            try:
                page_text = page.get_text("text") or ""
            except Exception:
                page_text = ""
            if page_text:
                parts.append(page_text)
                total += len(page_text)
                if total >= MAX_TEXT_CHARS:
                    break
        return "\n\n".join(parts)
    finally:
        doc.close()


def _extract_docx(data: bytes) -> str:
    import zipfile

    from docx import Document

    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            total_uncompressed = sum(entry.file_size for entry in archive.infolist())
            if total_uncompressed > 25 * 1024 * 1024:
                raise DocumentError("This DOCX file expands to an unsafe size.")
            for entry in archive.infolist():
                if entry.compress_size and entry.file_size / entry.compress_size > 200:
                    raise DocumentError("This DOCX file has an unsafe compression ratio.")
    except zipfile.BadZipFile as error:
        raise DocumentError("This DOCX file could not be opened.") from error

    try:
        document = Document(io.BytesIO(data))
    except Exception as error:
        raise DocumentError("This DOCX file could not be opened.") from error

    parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(dict.fromkeys(cells)))
    return "\n".join(parts)


def _extract_plain_text(data: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise DocumentError("This text file uses an unsupported encoding.")


def extract_document(filename: str, data: bytes) -> tuple[str, str, str]:
    """Returns (safe_filename, mime_type, extracted_text)."""
    safe_name = sanitize_filename(filename)
    extension = Path(safe_name).suffix.lower()
    if extension not in ALLOWED_EXTENSIONS:
        raise DocumentError("Unsupported file type. Upload a PDF, DOCX, TXT, or MD file.")
    if not data:
        raise DocumentError("The uploaded file is empty.")
    if len(data) > MAX_FILE_BYTES:
        raise DocumentError(f"Files must be {MAX_FILE_BYTES // (1024 * 1024)} MB or smaller.")

    try:
        if extension == ".pdf":
            text = _extract_pdf(data)
        elif extension == ".docx":
            text = _extract_docx(data)
        else:
            text = _extract_plain_text(data)
    except DocumentError:
        raise
    except Exception as error:
        raise DocumentError("The document could not be read. It may be damaged.") from error

    cleaned = _clean_text(text)
    if len(cleaned) < 20:
        raise DocumentError(
            "No readable text was found. Scanned PDFs and image-only files are not supported."
        )
    return safe_name, ALLOWED_EXTENSIONS[extension], cleaned


def _read_local() -> list[dict]:
    path = config.DOCUMENTS_FILE
    if not path.exists():
        return []
    try:
        with path.open("r", encoding="utf-8") as handle:
            data = json.load(handle)
        documents = data.get("documents") if isinstance(data, dict) else data
        return documents if isinstance(documents, list) else []
    except (OSError, json.JSONDecodeError):
        return []


def _write_local(documents: list[dict]) -> None:
    path = config.DOCUMENTS_FILE
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(path.suffix + ".tmp")
    with temporary.open("w", encoding="utf-8") as handle:
        json.dump({"documents": documents}, handle, ensure_ascii=False, indent=2)
    temporary.replace(path)


def _public(document: dict) -> dict:
    created = document.get("created_at")
    if isinstance(created, datetime):
        created = created.isoformat()
    return {
        "id": str(document["id"]),
        "owner_type": document["owner_type"],
        "owner_key": document["owner_key"],
        "filename": document["filename"],
        "mime_type": document.get("mime_type", ""),
        "size_bytes": int(document.get("size_bytes") or 0),
        "characters": len(document.get("content_text") or ""),
        "created_at": created or "",
    }


def count_documents(owner_type: str, owner_key: str) -> int:
    owner_type, owner_key = normalize_owner(owner_type, owner_key)
    if database.is_configured():
        with _get_connection() as connection:
            row = connection.execute(
                "SELECT COUNT(*) AS count FROM documents WHERE owner_type = %s AND owner_key = %s",
                (owner_type, owner_key),
            ).fetchone()
        return int(row["count"]) if row else 0
    return sum(
        1
        for document in _read_local()
        if document.get("owner_type") == owner_type and document.get("owner_key") == owner_key
    )


def add_document(owner_type: str, owner_key: str, filename: str, data: bytes) -> dict:
    owner_type, owner_key = normalize_owner(owner_type, owner_key)
    if count_documents(owner_type, owner_key) >= MAX_DOCUMENTS_PER_OWNER:
        raise DocumentError(
            f"A profile can have at most {MAX_DOCUMENTS_PER_OWNER} documents. "
            "Delete one before uploading another."
        )

    safe_name, mime_type, text = extract_document(filename, data)
    document = {
        "id": uuid.uuid4().hex,
        "owner_type": owner_type,
        "owner_key": owner_key,
        "filename": safe_name,
        "mime_type": mime_type,
        "size_bytes": len(data),
        "content_text": text,
        "created_at": datetime.now(timezone.utc).isoformat(),
    }

    if database.is_configured():
        with _get_connection() as connection:
            connection.execute(
                """
                INSERT INTO documents
                    (id, owner_type, owner_key, filename, mime_type,
                     size_bytes, content_text)
                VALUES (%(id)s, %(owner_type)s, %(owner_key)s, %(filename)s,
                        %(mime_type)s, %(size_bytes)s, %(content_text)s)
                """,
                document,
            )
        return _public(document)

    with _FILE_LOCK:
        documents = _read_local()
        documents.append(document)
        _write_local(documents)
    return _public(document)


def list_documents(owner_type: str, owner_key: str) -> list[dict]:
    owner_type, owner_key = normalize_owner(owner_type, owner_key)
    if database.is_configured():
        with _get_connection() as connection:
            rows = connection.execute(
                """
                SELECT id, owner_type, owner_key, filename, mime_type,
                       size_bytes, content_text, created_at
                FROM documents
                WHERE owner_type = %s AND owner_key = %s
                ORDER BY created_at, id
                """,
                (owner_type, owner_key),
            ).fetchall()
        return [_public(dict(row)) for row in rows]

    return [
        _public(document)
        for document in _read_local()
        if document.get("owner_type") == owner_type and document.get("owner_key") == owner_key
    ]


def delete_document(document_id: str) -> bool:
    document_id = (document_id or "").strip()
    if not document_id or len(document_id) > 64:
        return False
    if database.is_configured():
        with _get_connection() as connection:
            cursor = connection.execute("DELETE FROM documents WHERE id = %s", (document_id,))
            return cursor.rowcount == 1

    with _FILE_LOCK:
        documents = _read_local()
        remaining = [document for document in documents if document.get("id") != document_id]
        if len(remaining) == len(documents):
            return False
        _write_local(remaining)
    return True


def delete_documents_for(owner_type: str, owner_key: str) -> None:
    owner_type, owner_key = normalize_owner(owner_type, owner_key)
    if database.is_configured():
        with _get_connection() as connection:
            connection.execute(
                "DELETE FROM documents WHERE owner_type = %s AND owner_key = %s",
                (owner_type, owner_key),
            )
        return

    with _FILE_LOCK:
        documents = _read_local()
        remaining = [
            document
            for document in documents
            if not (
                document.get("owner_type") == owner_type
                and document.get("owner_key") == owner_key
            )
        ]
        if len(remaining) != len(documents):
            _write_local(remaining)


def move_documents(owner_type: str, old_key: str, new_key: str) -> None:
    owner_type, old_key = normalize_owner(owner_type, old_key)
    _, new_key = normalize_owner(owner_type, new_key)
    if old_key == new_key:
        return
    if database.is_configured():
        with _get_connection() as connection:
            connection.execute(
                "UPDATE documents SET owner_key = %s WHERE owner_type = %s AND owner_key = %s",
                (new_key, owner_type, old_key),
            )
        return

    with _FILE_LOCK:
        documents = _read_local()
        changed = False
        for document in documents:
            if document.get("owner_type") == owner_type and document.get("owner_key") == old_key:
                document["owner_key"] = new_key
                changed = True
        if changed:
            _write_local(documents)


def _load_texts(owner_type: str, owner_key: str) -> list[dict]:
    owner_type, owner_key = normalize_owner(owner_type, owner_key)
    if database.is_configured():
        with _get_connection() as connection:
            rows = connection.execute(
                """
                SELECT filename, content_text
                FROM documents
                WHERE owner_type = %s AND owner_key = %s
                ORDER BY created_at, id
                """,
                (owner_type, owner_key),
            ).fetchall()
        return [dict(row) for row in rows]

    return [
        {"filename": document.get("filename", "document"), "content_text": document.get("content_text", "")}
        for document in _read_local()
        if document.get("owner_type") == owner_type and document.get("owner_key") == owner_key
    ]


def build_reference_text(owner_type: str, owner_key: str, limit: int = MAX_PROMPT_CHARS) -> str:
    """Returns prompt-ready reference text, safely delimited and length-capped."""
    try:
        documents = _load_texts(owner_type, owner_key)
    except Exception:
        return ""
    if not documents:
        return ""

    parts: list[str] = []
    remaining = max(1000, limit)
    for document in documents:
        text = (document.get("content_text") or "").strip()
        if not text or remaining <= 0:
            continue
        header = f"[Document: {document.get('filename', 'document')}]\n"
        chunk = text[: max(0, remaining - len(header))]
        # Prevent the model from closing our untrusted-data fence early.
        chunk = chunk.replace("<<<REFERENCE_END>>>", "").replace("<<<REFERENCE_START>>>", "")
        parts.append(header + chunk)
        remaining -= len(header) + len(chunk) + 2

    return "\n\n".join(parts).strip()
